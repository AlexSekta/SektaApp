import os
import time

from collections import defaultdict
from datetime import datetime
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy.orm import Session

from backend.logging_config import get_logger
from backend.models.master_product import MasterProduct
from backend.models.one_c_stock import OneCStock

logger = get_logger("one_c_stock_import")

DEFAULT_EXCHANGE_DIR = r"C:\Users\Евгения\WB_Analytics\1c_exchange"

ONE_C_EXCHANGE_DIR = Path(
    os.getenv(
        "ONE_C_EXCHANGE_DIR",
        DEFAULT_EXCHANGE_DIR,
    )
)

FILE_CONFIG = {
    "1.xlsx": "Маркетплейс",
    "2.xlsx": "Готовая продукция",
}


class OneCImportError(Exception):
    pass


def _normalize_header(value):

    if value is None:
        return ""

    return str(value).strip().lower()


def _parse_stock_value(
    value,
    *,
    file_name: str,
    row_number: int,
):

    if value is None:
        raise OneCImportError(f"{file_name}: строка {row_number}: " "не указан остаток")

    if isinstance(value, bool):
        raise OneCImportError(
            f"{file_name}: строка {row_number}: " "некорректный остаток"
        )

    if isinstance(value, int):
        stock = value

    elif isinstance(value, float):

        if not value.is_integer():
            raise OneCImportError(
                f"{file_name}: строка {row_number}: "
                f"остаток должен быть целым числом, получено {value}"
            )

        stock = int(value)

    else:

        text = str(value).strip()

        try:
            stock = int(text)

        except ValueError as error:
            raise OneCImportError(
                f"{file_name}: строка {row_number}: " f"некорректный остаток '{text}'"
            ) from error

    if stock < 0:
        raise OneCImportError(
            f"{file_name}: строка {row_number}: " "остаток не может быть отрицательным"
        )

    return stock


def _ensure_file_stable(
    path: Path,
    wait_seconds: float = 1.0,
):

    first = path.stat()

    time.sleep(wait_seconds)

    second = path.stat()

    if first.st_size != second.st_size or first.st_mtime_ns != second.st_mtime_ns:
        raise OneCImportError(
            f"{path.name}: файл ещё записывается. " "Импорт остановлен."
        )


def _read_stock_file(path: Path):

    try:
        workbook = load_workbook(
            filename=path,
            read_only=True,
            data_only=True,
        )

    except Exception as error:
        raise OneCImportError(f"{path.name}: не удалось открыть Excel-файл") from error

    try:

        if not workbook.worksheets:
            raise OneCImportError(f"{path.name}: в файле нет листов")

        sheet = workbook.worksheets[0]

        header_map = {}

        for column_index in range(
            1,
            sheet.max_column + 1,
        ):
            header = _normalize_header(
                sheet.cell(
                    row=1,
                    column=column_index,
                ).value
            )

            if header:
                header_map[header] = column_index

        article_column = header_map.get("артикул")

        stock_column = header_map.get("остатки")

        if not article_column:
            raise OneCImportError(f"{path.name}: не найдена колонка 'Артикул'")

        if not stock_column:
            raise OneCImportError(f"{path.name}: не найдена колонка 'Остатки'")

        stocks = defaultdict(int)
        duplicate_rows = defaultdict(list)
        blank_article_rows = []

        for row_number in range(
            2,
            sheet.max_row + 1,
        ):
            article_value = sheet.cell(
                row=row_number,
                column=article_column,
            ).value

            stock_value = sheet.cell(
                row=row_number,
                column=stock_column,
            ).value

            article = str(article_value).strip() if article_value is not None else ""

            if not article and stock_value is None:
                continue

            # В реальном 2.xlsx есть строки с остатком,
            # но без артикула. Сопоставить их невозможно.
            if not article:
                blank_article_rows.append(row_number)
                continue

            stock = _parse_stock_value(
                stock_value,
                file_name=path.name,
                row_number=row_number,
            )

            if article in stocks:
                duplicate_rows[article].append(row_number)

            # Дубли одного артикула агрегируем суммой.
            stocks[article] += stock

        if not stocks:
            raise OneCImportError(
                f"{path.name}: файл не содержит " "ни одного товара с артикулом"
            )

        return {
            "stocks": dict(stocks),
            "blank_article_rows": blank_article_rows,
            "duplicates": {article: rows for article, rows in duplicate_rows.items()},
            "sheet": sheet.title,
        }

    finally:
        workbook.close()


def _validate_master_products(
    db: Session,
    parsed_files: dict,
):

    missing_by_group = {}

    for file_name, group in FILE_CONFIG.items():

        master_articles = {
            str(item.seller_article).strip()
            for item in (
                db.query(MasterProduct)
                .filter(
                    MasterProduct.active == True,
                    MasterProduct.group == group,
                )
                .all()
            )
            if item.seller_article
        }

        imported_articles = set(parsed_files[file_name]["stocks"].keys())

        missing = sorted(master_articles - imported_articles)

        if missing:
            missing_by_group[group] = missing

    if missing_by_group:

        details = []

        for group, articles in missing_by_group.items():
            preview = ", ".join(articles[:10])

            if len(articles) > 10:
                preview += f" и ещё {len(articles) - 10}"

            details.append(f"{group}: {preview}")

        raise OneCImportError(
            "В файлах 1С отсутствуют товары "
            "из реестра. Импорт остановлен. " + " | ".join(details)
        )


def import_one_c_stocks(
    db: Session,
):

    logger.info(
        "1C_IMPORT_START | path=%s",
        ONE_C_EXCHANGE_DIR,
    )

    if not ONE_C_EXCHANGE_DIR.exists():

        logger.error(
            "1C_IMPORT_FAILED | " "reason=directory_missing | " "path=%s",
            ONE_C_EXCHANGE_DIR,
        )

        raise OneCImportError("Папка обмена 1С не найдена: " f"{ONE_C_EXCHANGE_DIR}")

    paths = {file_name: ONE_C_EXCHANGE_DIR / file_name for file_name in FILE_CONFIG}

    missing_files = [
        file_name for file_name, path in paths.items() if not path.is_file()
    ]

    if missing_files:

        logger.error(
            "1C_IMPORT_FAILED | " "reason=file_missing | " "files=%s",
            ",".join(missing_files),
        )

        raise OneCImportError(
            "Не найдены обязательные файлы 1С: "
            + ", ".join(missing_files)
            + ". Импорт остановлен."
        )

    # Не читаем Excel, пока внешний процесс
    # ещё может дописывать файл.
    for path in paths.values():
        _ensure_file_stable(path)

    parsed_files = {}

    for file_name, path in paths.items():

        parsed = _read_stock_file(path)

        parsed_files[file_name] = parsed

        logger.info(
            "1C_FILE_VALIDATED | "
            "file=%s | "
            "rows=%s | "
            "blank_articles=%s | "
            "duplicates=%s",
            file_name,
            len(parsed["stocks"]),
            len(parsed["blank_article_rows"]),
            len(parsed["duplicates"]),
        )

    # Каждый активный товар реестра обязан
    # присутствовать в своём файле 1С.
    _validate_master_products(
        db,
        parsed_files,
    )

    imported_at = datetime.utcnow()

    try:

        # Оба склада заменяются в ОДНОЙ транзакции.
        for file_name, group in FILE_CONFIG.items():

            db.query(OneCStock).filter(OneCStock.group == group).delete(
                synchronize_session=False
            )

            for article, stock in parsed_files[file_name]["stocks"].items():
                db.add(
                    OneCStock(
                        seller_article=article,
                        group=group,
                        stock=stock,
                        source_file=file_name,
                        updated_at=imported_at,
                    )
                )

        db.commit()

    except Exception:

        db.rollback()

        logger.exception("1C_IMPORT_FAILED | " "reason=database_error")

        raise

    # Удаляем файлы только ПОСЛЕ commit.
    deleted_files = []

    try:

        for file_name, path in paths.items():

            path.unlink()

            deleted_files.append(file_name)

    except Exception as error:

        logger.exception(
            "1C_FILES_DELETE_FAILED | " "deleted=%s",
            ",".join(deleted_files),
        )

        raise OneCImportError(
            "Данные 1С записаны в БД, "
            "но не удалось удалить все исходные файлы. "
            "Автоматический процесс нужно остановить "
            "и проверить папку обмена."
        ) from error

    total_saved = sum(
        len(parsed_files[file_name]["stocks"]) for file_name in FILE_CONFIG
    )

    result = {
        "status": "ok",
        "saved": total_saved,
        "files": {},
    }

    for file_name, group in FILE_CONFIG.items():

        parsed = parsed_files[file_name]

        result["files"][file_name] = {
            "group": group,
            "saved": len(parsed["stocks"]),
            "blank_article_rows": len(parsed["blank_article_rows"]),
            "duplicate_articles": sorted(parsed["duplicates"].keys()),
        }

    logger.info(
        "1C_IMPORT_OK | " "saved=%s | " "files_deleted=%s",
        total_saved,
        ",".join(deleted_files),
    )

    return result
