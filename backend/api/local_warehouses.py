"""Собственные склады и безопасная загрузка Excel; не меняет синхронизацию МП."""
from datetime import datetime
from io import BytesIO
from fastapi import APIRouter, Depends, HTTPException, UploadFile, File
from fastapi.responses import StreamingResponse
from pydantic import BaseModel
from sqlalchemy import text, func
from sqlalchemy.orm import Session
from openpyxl import Workbook, load_workbook
from backend.dependencies import get_db
from backend.models.master_product import MasterProduct
from backend.models.one_c_stock import OneCStock

router = APIRouter(prefix='/local-warehouses', tags=['Собственные склады'])
DEFAULTS = ('Маркетплейс', 'Готовая продукция')

class WarehouseCreate(BaseModel):
    name: str


def org(db):
    value = db.execute(text("SELECT NULLIF(current_setting('app.organization_id', true), '')::integer")).scalar()
    if value is None:
        raise HTTPException(401, 'Нет контекста организации')
    return value


def warehouses(db):
    oid=org(db)
    return [r[0] for r in db.execute(
        text('SELECT name FROM local_warehouses WHERE organization_id=:o ORDER BY id'), {'o':oid})]

@router.get('')
def list_warehouses(db: Session=Depends(get_db)):
    return [{'name': name, 'can_delete': True} for name in warehouses(db)]

@router.post('')
def create_warehouse(item: WarehouseCreate, db: Session=Depends(get_db)):
    name=item.name.strip()
    if not 1 <= len(name) <= 100 or any(c in name for c in '\\/:*?[]'):
        raise HTTPException(422,'Некорректное название склада')
    oid=org(db)
    if name.lower() in {x.lower() for x in warehouses(db)}:
        raise HTTPException(409,'Склад с таким названием уже существует')
    db.execute(text('INSERT INTO local_warehouses(organization_id,name) VALUES (:o,:n)'),{'o':oid,'n':name})
    db.commit()
    return {'status':'ok','name':name}

@router.delete('/{name}')
def delete_warehouse(name:str, db:Session=Depends(get_db)):
    oid=org(db)
    if name not in warehouses(db):
        raise HTTPException(404, 'Склад не найден')
    if db.query(MasterProduct).filter(
        MasterProduct.organization_id == oid, MasterProduct.group == name
    ).first():
        raise HTTPException(409, 'На складе есть привязанные товары. Сначала перенесите их на другой склад.')
    if db.query(OneCStock).filter(
        OneCStock.organization_id == oid, OneCStock.group == name,
        OneCStock.stock != 0
    ).first():
        raise HTTPException(409, 'На складе есть ненулевые остатки. Удаление запрещено.')
    # Нулевые записи не являются запасом; удаляем вместе с пустым складом.
    db.query(OneCStock).filter(
        OneCStock.organization_id == oid, OneCStock.group == name,
        OneCStock.stock == 0
    ).delete(synchronize_session=False)
    db.execute(text('DELETE FROM local_warehouses WHERE organization_id=:o AND name=:n'),
               {'o':oid,'n':name})
    db.commit()
    return {'status':'ok'}

class WarehouseProductCreate(BaseModel):
    article: str
    group: str


@router.post('/products')
def add_product(item: WarehouseProductCreate, db: Session = Depends(get_db)):
    """Один артикул принадлежит только одному складу организации."""
    oid = org(db)
    article = item.article.strip()
    group = item.group.strip()
    if not article or len(article) > 128:
        raise HTTPException(422, 'Введите корректный артикул')
    if group not in warehouses(db):
        raise HTTPException(422, 'Склад не найден')
    existing = db.query(MasterProduct).filter(
        MasterProduct.organization_id == oid,
        MasterProduct.seller_article == article,
    ).first()
    if existing:
        if existing.group == group:
            raise HTTPException(409, f'Артикул {article} уже находится на складе «{group}»')
        raise HTTPException(409, f'Артикул {article} уже привязан к складу «{existing.group}». Товар может находиться только на одном складе.')
    internal = db.query(MasterProduct).filter(
        MasterProduct.organization_id == oid,
        MasterProduct.internal_article == article,
    ).first()
    if internal:
        raise HTTPException(409, 'Такой внутренний артикул уже существует')
    order = db.query(func.max(MasterProduct.sort_order)).filter(MasterProduct.organization_id == oid).scalar() or 0
    db.add(MasterProduct(
        organization_id=oid, seller_article=article, internal_article=article,
        name=article, group=group, active=True, sort_order=order + 1,
        status='READY', purchase_price=0, additional_costs=0,
    ))
    db.commit()
    return {'status': 'ok', 'article': article, 'group': group}


@router.get('/template')
def template(db:Session=Depends(get_db)):
    w=Workbook();sheet=w.active;sheet.title='Остатки';sheet.append(['Склад','Артикул','Количество'])
    valid=set(warehouses(db))
    for p in db.query(MasterProduct).filter(MasterProduct.active.is_(True)).order_by(MasterProduct.group,MasterProduct.seller_article):
        if p.group not in valid:continue
        stock=db.query(OneCStock).filter(OneCStock.group==p.group,OneCStock.seller_article==p.seller_article).first()
        sheet.append([p.group,p.seller_article,stock.stock if stock else 0])
    sheet.column_dimensions['A'].width=28;sheet.column_dimensions['B'].width=24;sheet.column_dimensions['C'].width=18
    buf=BytesIO();w.save(buf);buf.seek(0)
    return StreamingResponse(buf,media_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet',headers={'Content-Disposition':'attachment; filename="stock_template.xlsx"'})

@router.post('/import')
async def import_file(file:UploadFile=File(...),db:Session=Depends(get_db)):
    if not (file.filename or '').lower().endswith('.xlsx'):
        raise HTTPException(422,'Поддерживается только .xlsx')
    contents=await file.read(5_000_001)
    if len(contents)>5_000_000:
        raise HTTPException(413,'Файл превышает 5 МБ')
    try:
        wb=load_workbook(BytesIO(contents),read_only=True,data_only=True)
    except Exception:
        raise HTTPException(422,'Не удалось прочитать Excel')
    try:
        sheet=wb.active
        rows=sheet.iter_rows(values_only=True)
        headers=[str(x or '').strip().lower() for x in next(rows,())]
        if headers[:3]!=['склад','артикул','количество']:
            raise HTTPException(422,'Ожидаются колонки: Склад | Артикул | Количество')
        products={p.seller_article:p.group for p in db.query(MasterProduct).filter(MasterProduct.active.is_(True))}
        wh=set(warehouses(db));updates={};unknown=[]
        for number,row in enumerate(rows,2):
            values=list(row)+[None,None,None]
            warehouse=str(values[0] or '').strip();article=str(values[1] or '').strip();qty=values[2]
            if not warehouse and not article and qty is None:continue
            if not warehouse or not article or warehouse not in wh or article not in products:
                unknown.append(f'{number}: склад «{warehouse}», артикул {article} — неизвестный склад или артикул');continue
            assigned_group = products[article]
            if assigned_group != warehouse:
                unknown.append(f'{number}: артикул {article} закреплён за складом «{assigned_group}», а в файле указан «{warehouse}»');continue
            if (warehouse,article) in updates:
                raise HTTPException(422,f'Дубликат артикула {article} на складе {warehouse}, строка {number}')
            if qty is None or str(qty).strip().upper() in ('','#Н/Д','#N/A'):
                count=0
            else:
                if isinstance(qty,bool):raise HTTPException(422,f'Некорректное количество в строке {number}')
                try:
                    count=int(qty)
                    if str(qty).strip() not in (str(count),str(float(count))):raise ValueError()
                except (ValueError,TypeError):
                    raise HTTPException(422,f'Некорректное количество в строке {number}')
                if count<0:raise HTTPException(422,f'Отрицательное количество в строке {number}')
            updates[(warehouse,article)]=count
        if not updates:
            raise HTTPException(422,'Нет известных артикулов для обновления')
        for (warehouse,article),count in updates.items():
            record=db.query(OneCStock).filter(OneCStock.group==warehouse,OneCStock.seller_article==article).first()
            if record:
                record.stock=count;record.source_file='Загрузка через интерфейс';record.updated_at=datetime.utcnow()
            else:
                db.add(OneCStock(organization_id=org(db),group=warehouse,seller_article=article,stock=count,source_file='Загрузка через интерфейс'))
        db.commit()
        return {'status':'ok','updated':len(updates),'unknown':unknown[:100], 'unknown_count':len(unknown)}
    except HTTPException:
        db.rollback();raise
    except Exception:
        db.rollback();raise
    finally:
        wb.close()


class ManualStockUpdate(BaseModel):
    article: str
    group: str
    stock: int
    expected_stock: int


@router.put('/stock')
def set_local_stock(item: ManualStockUpdate, db: Session = Depends(get_db)):
    """Точечная установка остатка 1С; не отправляет данные в маркетплейсы."""
    oid = org(db)
    if item.stock < 0 or item.expected_stock < 0:
        raise HTTPException(422, 'Количество не может быть отрицательным')
    if item.stock > 2_147_483_647:
        raise HTTPException(422, 'Слишком большое количество')
    product = db.query(MasterProduct).filter(
        MasterProduct.organization_id == oid,
        MasterProduct.seller_article == item.article,
        MasterProduct.group == item.group,
        MasterProduct.active.is_(True),
    ).first()
    if not product:
        raise HTTPException(404, 'Товар не найден на указанном складе')
    # Блокировка товара защищает от одновременных ручных изменений.
    db.query(MasterProduct).filter(MasterProduct.id == product.id).with_for_update().one()
    record = db.query(OneCStock).filter(
        OneCStock.organization_id == oid,
        OneCStock.seller_article == item.article,
        OneCStock.group == item.group,
    ).with_for_update().first()
    current = record.stock if record else 0
    if current != item.expected_stock:
        db.rollback()
        raise HTTPException(409, f'Остаток уже изменился: сейчас {current} шт. Обновите страницу.')
    if record:
        record.stock = item.stock
        record.source_file = 'Ручное изменение в приложении'
        record.updated_at = datetime.utcnow()
    else:
        db.add(OneCStock(organization_id=oid, seller_article=item.article,
                         group=item.group, stock=item.stock,
                         source_file='Ручное изменение в приложении',
                         updated_at=datetime.utcnow()))
    db.commit()
    return {'status': 'ok', 'article': item.article, 'group': item.group,
            'previous_stock': current, 'stock': item.stock}
